"""
build_table1_main_stigma_0622.py
Produces Table 1 (main stigma scores) for the NEJM AI 14-model submission.

This is the 14-model sibling of scripts_public/build_table1_0529.py. The Word
table layout (2 panels, family grouping, 9 columns) is kept identical to the
accepted proof of an earlier submission (Table 1);
the only substantive changes are the 14th model (Gemini-3.5-Flash, Google
family) and the title-case caption "14 Large Language Models".

Inputs (14-model dataset of record):
  analysis_14models_20260614/en/table2_en_stigma_scores.csv
  analysis_14models_20260614/zh/table3_zh_stigma_scores.csv
    columns: model_label, Scale, PLE, DEP, APS, SCZ  (cells "39.5 [39.2, 39.9]")

Outputs (outputs/figures_tables/):
  table1_main_stigma_scores.docx   — Word table, Panel A (EN) over Panel B (ZH)
  table1_main_stigma_scores.csv    — verification CSV (long, hyphenated names)

Model names in both outputs follow scripts_public/Models_name.md (canonical
hyphenated display names; raw data ids in the CSV are spaced).

Usage:
  python3 scripts_public/build_table1_main_stigma_0622.py
"""

import csv
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).parent.parent
EN_CSV = ROOT / "analysis_14models_20260614/en/table2_en_stigma_scores.csv"
ZH_CSV = ROOT / "analysis_14models_20260614/zh/table3_zh_stigma_scores.csv"
OUT_DIR = ROOT / "outputs" / "figures_tables"

N_MODELS = 14

# ── Model order & family grouping ─────────────────────────────────────────────
# Keys are the raw model_label values in the 14-model CSVs (spaced names with a
# deployment tag). Order = developer family (OpenAI → Anthropic → Google → Meta
# → DeepSeek), each family by release date, matching the proof Table 1.
# Gemini-3.5-Flash is the 14th model, inserted into the Google family after the
# Gemini-3.1 line and before Gemma (Gemma is open-weight and stays last in Google).
FAMILY_ORDER = [
    ("OpenAI", [
        "GPT-3.5 (API)",
        "GPT-4o (API)",
        "GPT-OSS-120B (Open-source)",
        "GPT-5.4 (API)",
        "GPT-5.4 Mini (API)",
    ]),
    ("Anthropic", [
        "Claude Sonnet 4.5 (API)",
        "Claude Opus 4.6 (API)",
    ]),
    ("Google", [
        "Gemini 2.5 Flash (API)",
        "Gemini 3.1 Pro Preview (API)",
        "Gemini 3.1 Flash-Lite Preview (API)",
        "Gemini 3.5 Flash (API)",
        "Gemma 4 31B IT (Open-source)",
    ]),
    ("Meta", [
        "Llama 3.3 70B (Open-source)",
    ]),
    ("DeepSeek", [
        "DeepSeek V3.2 (Open-source)",
    ]),
]

VIGNETTES = ["PLE", "DEP", "APS", "SCZ"]

# ── Display labels (canonical, per scripts_public/Models_name.md) ──────────────
DISPLAY_LABELS = {
    "GPT-3.5 (API)":                        "GPT-3.5",
    "GPT-4o (API)":                         "GPT-4o",
    "GPT-OSS-120B (Open-source)":           "GPT-OSS-120B",
    "GPT-5.4 (API)":                        "GPT-5.4",
    "GPT-5.4 Mini (API)":                   "GPT-5.4-Mini",
    "Claude Sonnet 4.5 (API)":              "Claude Sonnet 4.5",
    "Claude Opus 4.6 (API)":                "Claude Opus 4.6",
    "Gemini 2.5 Flash (API)":               "Gemini-2.5-Flash",
    "Gemini 3.1 Pro Preview (API)":         "Gemini-3.1-Pro-Preview",
    "Gemini 3.1 Flash-Lite Preview (API)":  "Gemini-3.1-Flash-Lite-Preview",
    "Gemini 3.5 Flash (API)":               "Gemini-3.5-Flash",
    "Gemma 4 31B IT (Open-source)":         "Gemma-4-31B-IT",
    "Llama 3.3 70B (Open-source)":          "Llama-3.3-70B",
    "DeepSeek V3.2 (Open-source)":          "DeepSeek-V3.2",
}

# ── Shading ───────────────────────────────────────────────────────────────────
FAMILY_GRAY = "D9D9D9"   # light gray for header / family rows
WHITE       = "FFFFFF"


# ── Parse CSV into nested dict ────────────────────────────────────────────────
def fmt_cell(raw: str) -> str:
    """Convert '39.5 [39.2, 39.9]' → '39.5 (39.2, 39.9)'."""
    return re.sub(r"\[(.+?)\]", r"(\1)", raw.strip())


def load_data(path: Path) -> dict:
    """Returns {model_label: {'social_distance': {vig: str}, 'prejudice': {vig: str}}}."""
    data: dict = {}
    with open(path) as f:
        for row in csv.DictReader(f):
            ml = row["model_label"]
            scale = row["Scale"].lower().replace(" ", "_")
            data.setdefault(ml, {})[scale] = {v: fmt_cell(row[v]) for v in VIGNETTES}
    return data


def build_rows(data: dict) -> list:
    """Flatten FAMILY_ORDER into family-header and data rows for one panel."""
    rows = []
    for family_name, models in FAMILY_ORDER:
        rows.append({"type": "family", "name": family_name})
        for ml in models:
            if ml not in data:
                continue
            sd = data[ml].get("social_distance", {v: "—" for v in VIGNETTES})
            pj = data[ml].get("prejudice",       {v: "—" for v in VIGNETTES})
            rows.append({
                "type":  "data",
                "label": DISPLAY_LABELS.get(ml, ml),
                "sd":    sd,
                "pj":    pj,
            })
    return rows


# ── Word helpers ──────────────────────────────────────────────────────────────
def _set_cell_bg(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_para(cell, text: str, bold=False, size_pt=9,
               align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    return para


def _set_col_widths(table, widths_inches: list):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[i])


def _thick_border(cell, sides=("top", "bottom")):
    """Apply 1.5pt solid black border to specified sides of a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "12")   # 12 eighths of a point = 1.5pt
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_panel(doc: Document, panel_title: str, rows: list, is_first: bool = False):
    """Add one language panel (EN or ZH) to the document."""
    if not is_first:
        doc.add_paragraph()  # spacing between panels
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(panel_title)
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(10)

    # 9 columns: Model | SD×4 | PJ×4
    COL_W = [1.6] + [0.615] * 8
    NCOLS = 9

    table = doc.add_table(rows=0, cols=NCOLS)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"

    # ── Header row 1: scale group labels ──────────────────────────────────────
    hr1 = table.add_row()
    hr1.cells[0].text = ""
    sd_cell = hr1.cells[1].merge(hr1.cells[4])
    _cell_para(sd_cell, "Social Distance Score", bold=True, size_pt=9,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    pj_cell = hr1.cells[5].merge(hr1.cells[8])
    _cell_para(pj_cell, "Prejudice Score", bold=True, size_pt=9,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in hr1.cells:
        _set_cell_bg(cell, FAMILY_GRAY)
        _thick_border(cell, sides=["top"])

    # ── Header row 2: vignette labels ─────────────────────────────────────────
    hr2 = table.add_row()
    _cell_para(hr2.cells[0], "Model", bold=True, size_pt=9)
    _set_cell_bg(hr2.cells[0], FAMILY_GRAY)
    for i, vig in enumerate(VIGNETTES):
        _cell_para(hr2.cells[1 + i], vig, bold=True, size_pt=9,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_para(hr2.cells[5 + i], vig, bold=True, size_pt=9,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(hr2.cells[1 + i], FAMILY_GRAY)
        _set_cell_bg(hr2.cells[5 + i], FAMILY_GRAY)
    for cell in hr2.cells:
        _thick_border(cell, sides=["bottom"])

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_info in rows:
        if row_info["type"] == "family":
            dr = table.add_row()
            merged = dr.cells[0]
            for ci in range(1, NCOLS):
                merged = merged.merge(dr.cells[ci])
            _cell_para(merged, row_info["name"], bold=True, size_pt=9)
            _set_cell_bg(merged, FAMILY_GRAY)
        else:
            dr = table.add_row()
            _cell_para(dr.cells[0], "  " + row_info["label"], bold=False, size_pt=9)
            for i, vig in enumerate(VIGNETTES):
                _cell_para(dr.cells[1 + i], row_info["sd"].get(vig, "—"),
                           size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _cell_para(dr.cells[5 + i], row_info["pj"].get(vig, "—"),
                           size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            for cell in dr.cells:
                _set_cell_bg(cell, WHITE)

    for cell in table.rows[-1].cells:
        _thick_border(cell, sides=["bottom"])

    _set_col_widths(table, COL_W)


# ── Build verification CSV ────────────────────────────────────────────────────
def build_csv(en_data: dict, zh_data: dict, out_path: Path):
    fieldnames = ["language", "family", "model"] + \
                 [f"SD_{v}" for v in VIGNETTES] + \
                 [f"PJ_{v}" for v in VIGNETTES]
    with open(out_path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for lang, data in [("EN", en_data), ("ZH", zh_data)]:
            for family_name, models in FAMILY_ORDER:
                for ml in models:
                    if ml not in data:
                        continue
                    sd = data[ml].get("social_distance", {})
                    pj = data[ml].get("prejudice", {})
                    row = {
                        "language": lang,
                        "family":   family_name,
                        "model":    DISPLAY_LABELS.get(ml, ml),
                    }
                    for v in VIGNETTES:
                        row[f"SD_{v}"] = sd.get(v, "—")
                        row[f"PJ_{v}"] = pj.get(v, "—")
                    w.writerow(row)


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    en_data = load_data(EN_CSV)
    zh_data = load_data(ZH_CSV)

    en_rows = build_rows(en_data)
    zh_rows = build_rows(zh_data)

    doc = Document()
    section = doc.sections[0]
    section.page_width    = Inches(8.5)
    section.page_height   = Inches(11)
    section.left_margin   = Inches(0.75)
    section.right_margin  = Inches(0.75)
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    caption = doc.add_paragraph()
    caption_run = caption.add_run(
        f"Table 1. Mean (95% CI) Stigma Scores Across {N_MODELS} Large Language "
        "Models, Four Vignettes, and Two Language Conditions"
    )
    caption_run.bold = True
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10)

    add_panel(doc, "Panel A: English Prompts",          en_rows, is_first=True)
    add_panel(doc, "Panel B: Mandarin Chinese Prompts", zh_rows, is_first=False)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Values are mean (95% CI) across qualifying simulation runs "
        "(runs with ≥80% item completeness). "
        "SD = Social Distance scale (19 items; possible range 19–76); "
        "PJ = Prejudice scale (21 items; possible range 21–84). "
        "Vignette labels: PLE = psychotic-like experiences; "
        "DEP = major depression; APS = attenuated psychosis syndrome; "
        "SCZ = schizophrenia. Higher scores indicate greater stigma."
    )
    note_run.font.name = "Arial"
    note_run.font.size = Pt(8)
    note.paragraph_format.space_before = Pt(4)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docx_path = OUT_DIR / "table1_main_stigma_scores.docx"
    doc.save(docx_path)
    print(f"Saved: {docx_path}")

    csv_path = OUT_DIR / "table1_main_stigma_scores.csv"
    build_csv(en_data, zh_data, csv_path)
    print(f"Saved: {csv_path}")

    data_rows = sum(1 for r in en_rows if r["type"] == "data")
    family_rows = sum(1 for r in en_rows if r["type"] == "family")
    print(f"\nTable: 2 header rows + {family_rows} family rows + {data_rows} data "
          f"rows per panel × 2 panels (expected {N_MODELS} models per panel)")
    print("Columns: 9 (Model + SD×4 + PJ×4)")


if __name__ == "__main__":
    main()
